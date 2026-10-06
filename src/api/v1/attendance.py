import io
from typing import List, Optional
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, status, Query
from fastapi.responses import Response
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.core.database import get_db
from src.models.admin import Admin
from src.models.attendance import AttendanceLesson
from src.repositories.attendance_repo import AttendanceRepository
from src.schemas.attendance import (
    AttendanceLessonCreate,
    AttendanceLessonResponse,
    AttendanceSubmitRequest,
    LessonWithRosterResponse,
    ClubAttendanceResultsResponse,
)
from src.api.deps import get_current_admin, check_club_access

router = APIRouter(prefix="/attendance", tags=["Attendance & Results"])


@router.get("/lessons", response_model=List[AttendanceLessonResponse])
async def list_lessons(
    club_id: int = Query(..., description="To'garak ID"),
    current_admin: Admin = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    """Retrieve all lessons recorded for a club."""
    check_club_access(current_admin, club_id)
    repo = AttendanceRepository(session)
    return await repo.get_lessons_for_club(club_id)


@router.post("/lessons", response_model=AttendanceLessonResponse, status_code=status.HTTP_201_CREATED)
async def create_lesson(
    data: AttendanceLessonCreate,
    current_admin: Admin = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    """
    Create a new lesson and initialize attendance roster for all active students.
    (Requirement: O‘qituvchi har darsda davomat qiladi.)
    """
    check_club_access(current_admin, data.club_id)
    repo = AttendanceRepository(session)
    lesson = await repo.create_lesson_with_students(
        club_id=data.club_id,
        lesson_date=data.lesson_date,
        topic=data.topic,
        created_by_id=current_admin.id
    )
    # Fetch populated summary
    lessons = await repo.get_lessons_for_club(data.club_id)
    for l in lessons:
        if l["id"] == lesson.id:
            return l
    return AttendanceLessonResponse(
        id=lesson.id,
        club_id=lesson.club_id,
        lesson_date=lesson.lesson_date,
        topic=lesson.topic,
        created_at=lesson.created_at,
        present_count=0,
        absent_count=0,
        excused_count=0,
        total_records=0
    )


@router.get("/lessons/{lesson_id}", response_model=LessonWithRosterResponse)
async def get_lesson_roster(
    lesson_id: int,
    current_admin: Admin = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    """Get lesson attendance roster for marking students."""
    repo = AttendanceRepository(session)
    roster_data = await repo.get_lesson_roster(lesson_id)
    if not roster_data:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Dars topilmadi.")

    check_club_access(current_admin, roster_data["lesson"]["club_id"])
    return roster_data


@router.post("/lessons/{lesson_id}/save")
async def save_lesson_roster(
    lesson_id: int,
    data: AttendanceSubmitRequest,
    current_admin: Admin = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    """Save/update marked attendance for a lesson."""
    repo = AttendanceRepository(session)
    lesson = await repo.get_by_id(lesson_id)
    if not lesson:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Dars topilmadi.")

    check_club_access(current_admin, lesson.club_id)
    records_list = [r.model_dump() for r in data.records]
    success = await repo.save_roster(lesson_id, records_list)
    if not success:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Davomatni saqlab bo'lmadi.")

    return {"message": "Davomat muvaffaqiyatli saqlandi!"}


@router.delete("/lessons/{lesson_id}")
async def delete_lesson(
    lesson_id: int,
    current_admin: Admin = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    """Delete a lesson session."""
    repo = AttendanceRepository(session)
    lesson = await repo.get_by_id(lesson_id)
    if not lesson:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Dars topilmadi.")

    check_club_access(current_admin, lesson.club_id)
    await repo.delete(lesson_id)
    return {"message": "Dars jurnaldan o'chirildi."}


@router.get("/results", response_model=ClubAttendanceResultsResponse)
async def get_club_results(
    club_id: int = Query(..., description="To'garak ID"),
    current_admin: Admin = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    """
    Retrieve end-of-course attendance performance results and percentages for all enrolled students.
    (Requirement: Kurs oxirida tizim: nechta darsga qatnashgani, qatnashish foizi)
    """
    check_club_access(current_admin, club_id)
    repo = AttendanceRepository(session)
    results = await repo.get_club_attendance_results(club_id)
    if not results:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="To'garak topilmadi.")
    return results


@router.get("/export")
async def export_attendance_excel(
    club_id: int = Query(..., description="To'garak ID"),
    current_admin: Admin = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    """Generate and download Excel report for club attendance journal and percentage results."""
    check_club_access(current_admin, club_id)
    repo = AttendanceRepository(session)
    results = await repo.get_club_attendance_results(club_id)
    if not results:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="To'garak topilmadi.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Davomat va Natijalar"

    # Styling definitions
    title_font = Font(name="Calibri", size=15, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="4B5563")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    bold_data_font = Font(name="Calibri", size=11, bold=True)

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    success_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    warning_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    # Title Block
    ws.merge_cells("A1:G1")
    ws["A1"] = f"DAVOMAT VA NATIJALAR JURNALI: {results['club_name'].upper()}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:G2")
    ws["A2"] = f"O'tkazilgan jami darslar: {results['total_lessons_held']} ta | Sana: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Table Headers
    headers = [
        "№",
        "Talaba F.I.Sh",
        "Telefon Raqami",
        "Kursi",
        "Qatnashgan Darslari",
        "Qatnashish Foizi",
        "Bahosi / Natija"
    ]
    ws.append([]) # Row 3 empty
    ws.append(headers) # Row 4
    ws.row_dimensions[4].height = 26

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Data Rows
    row_idx = 5
    for idx, student in enumerate(results["students_summary"], start=1):
        pct = student["attendance_percentage"]
        row_data = [
            idx,
            student["full_name"],
            student["phone_number"] or "—",
            f"{student['course_level']}-kurs",
            f"{student['attended_lessons']} / {student['total_lessons']}",
            f"{pct}%",
            student["grade_label"]
        ]
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 22

        fill_to_use = alt_fill if row_idx % 2 == 0 else PatternFill(fill_type=None)
        if pct >= 90:
            fill_to_use = success_fill
        elif pct < 60:
            fill_to_use = warning_fill

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.font = bold_data_font if col_num in [2, 6] else data_font
            if fill_to_use.fill_type:
                cell.fill = fill_to_use
            cell.border = thin_border
            if col_num in [1, 4, 5, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        row_idx += 1

    # Adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"davomat_{results['club_id']}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
